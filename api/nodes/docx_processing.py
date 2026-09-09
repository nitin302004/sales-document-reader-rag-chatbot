import json
import re

from docx import Document
from langchain_core.messages import AIMessage

from api.utils.state import State


def read_docx(file_path: str) -> str:
    """Extract text from paragraphs and tables in a DOCX file."""
    doc = Document(file_path)
    text_content = []

    # Extract paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            text_content.append(para.text.strip())

    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            row_text = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]

            if row_text:
                text_content.append(" | ".join(row_text))

    return "\n".join(text_content)


def extract_field(text: str, field_names):
    """Extract a text value following a field label."""
    if isinstance(field_names, str):
        field_names = [field_names]

    for field_name in field_names:
        pattern = rf"{re.escape(field_name)}\s*[:|]\s*(.+)"
        match = re.search(pattern, text, re.IGNORECASE)

        if match:
            return match.group(1).strip()

    return None


def extract_number(text: str, field_names):
    """Extract a numeric value following a field label."""
    if isinstance(field_names, str):
        field_names = [field_names]

    for field_name in field_names:
        pattern = rf"{re.escape(field_name)}\s*[:|]\s*\$?\s*([\d,]+(?:\.\d+)?)"
        match = re.search(pattern, text, re.IGNORECASE)

        if match:
            return match.group(1).replace(",", "")

    return None


def extract_percentage(text: str, field_names):
    """Extract a percentage value following a field label."""
    if isinstance(field_names, str):
        field_names = [field_names]

    for field_name in field_names:
        pattern = rf"{re.escape(field_name)}\s*[:|]\s*([\d.]+\s*%)"
        match = re.search(pattern, text, re.IGNORECASE)

        if match:
            return match.group(1).strip()

    return None


def extract_date(text: str, field_names):
    """Extract a date value following a field label."""
    if isinstance(field_names, str):
        field_names = [field_names]

    date_pattern = (
        r"(\d{1,2}\s+"
        r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|"
        r"January|February|March|April|May|June|July|August|"
        r"September|October|November|December)"
        r"\s+\d{4}"
        r"|\d{4}-\d{2}-\d{2}"
        r"|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})"
    )

    for field_name in field_names:
        pattern = rf"{re.escape(field_name)}\s*[:|]\s*{date_pattern}"
        match = re.search(pattern, text, re.IGNORECASE)

        if match:
            return match.group(1).strip()

    return None


def docx_processing(state: State) -> State:
    """
    Process a sales-related DOCX document using
    rule-based field extraction.
    """

    print("Started processing the sales DOCX file...")

    file_bytes = state["contents"]

    # Temporarily save the uploaded DOCX
    file_path = "temp_sales_document.docx"

    with open(file_path, "wb") as f:
        f.write(file_bytes)

    # Extract document text
    text = read_docx(file_path)

    # Extract sales-related fields
    entities = {
        "order_id": extract_field(
            text,
            ["Order ID", "Order Number", "Order No"]
        ),

        "order_date": extract_date(
            text,
            ["Order Date", "Purchase Date"]
        ),

        "product": extract_field(
            text,
            ["Product", "Product Name", "Item"]
        ),

        "category": extract_field(
            text,
            ["Category", "Product Category"]
        ),

        "seller": extract_field(
            text,
            ["Seller", "Seller Name", "Vendor"]
        ),

        "customer": extract_field(
            text,
            ["Customer", "Customer Name", "Buyer"]
        ),

        "quantity": extract_number(
            text,
            ["Quantity", "Units Sold", "Units"]
        ),

        "unit_price": extract_number(
            text,
            ["Unit Price", "Price Per Unit", "Selling Price"]
        ),

        "discount": extract_percentage(
            text,
            ["Discount", "Discount Rate"]
        ),

        "revenue": extract_number(
            text,
            ["Revenue", "Sales Revenue", "Total Sales"]
        ),

        "region": extract_field(
            text,
            ["Region", "Sales Region", "Market"]
        ),

        "sales_channel": extract_field(
            text,
            ["Sales Channel", "Channel", "Marketplace"]
        ),
    }

    result = {
        "file_type": "docx",
        "document_type": "sales",
        "entities": entities
    }

    json_result = json.dumps(result, indent=2)

    response_message = AIMessage(content=json_result)

    state["messages"].append(response_message)

    print("Finished processing the sales DOCX file...")

    return state
